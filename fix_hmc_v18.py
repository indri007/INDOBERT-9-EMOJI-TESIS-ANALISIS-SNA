from docx import Document

src = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V17.docx"
out = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V18.docx"

doc = Document(src)
changed = []

# Exact paragraph replacements based on current V17 paragraph positions.
replacements = {
    126: """While the executive head of state (@prabowo) accumulated the highest In-Degree (15) as a prominent target of public accountability, his Out-Degree remained zero ($C_{\\text{out}} = 0$). In the observed network, this pattern indicates that interactions directed toward the account were not accompanied by recorded outbound interactions from the account within the analyzed corpus.""",

    127: """In contrast, xAI's conversational agent (@grok) recorded the highest Out-Degree in the observed network ($C_{\\text{out}} = 42$). Users tagged @grok in interactions concerning topics such as meal pricing and nutritional or poisoning-related information. This position indicates that the AI-mediated account occupied a distinct information-seeking position within the observed interaction network.""",

    158: """An important empirical finding concerns the communicative position of the artificial intelligence account. In classical crisis communication models (Coombs, 2007), epistemic authority is generally associated with institutional leaders, official spokespersons, or other recognized information sources. In the observed network, however, @grok recorded an Out-Degree of 42, while @prabowo recorded an Out-Degree of zero ($C_{\\text{out}} = 0$). This difference indicates an asymmetry in observed outbound interaction involving institutional and AI-mediated accounts."""
}

for n, replacement in replacements.items():
    idx = n - 1

    if idx >= len(doc.paragraphs):
        print(f"WARNING: P{n} does not exist")
        continue

    old = doc.paragraphs[idx].text

    if old != replacement:
        doc.paragraphs[idx].text = replacement
        changed.append((n, old, replacement))

doc.save(out)

print("HMC V18 CREATED")
print("=" * 80)
print("INPUT :", src)
print("OUTPUT:", out)
print("CHANGED PARAGRAPHS:", len(changed))
print()

for n, old, new in changed:
    print(f"P{n} CHANGED")
    print("OLD:", old)
    print("NEW:", new)
    print("-" * 80)
