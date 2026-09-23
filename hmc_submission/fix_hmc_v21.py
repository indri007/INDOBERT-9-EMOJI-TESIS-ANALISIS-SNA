from docx import Document

src = "HMC_REVIEW_MANUSCRIPT.docx"
out = "HMC_REVIEW_MANUSCRIPT_V21.docx"

doc = Document(src)
changed = []

# Exact paragraph replacements
replacements = {
    12: """The conceptualization of the Phygital Gap (Kotler, Kartajaya, & Setiawan, 2023) originated within commercial marketing and consumer experience management, explaining how discrepancies between omni-channel promises and consumer experiences may affect brand perceptions and switching behavior. However, political communication and public administration scholarship has only limited empirical application of this paradigm to state-sponsored welfare programs. In public governance, discrepancies between digitally communicated expectations and reported physical experiences may affect perceptions of political trust and institutional legitimacy, with implications for deliberative democratic processes (Levi & Stoker, 2000; Habermas, 1989, 2006; Coombs, 2007).""",

    17: """4. The Emergent Algorithmic Governance Gap (AI-Mediated Information Seeking vs. Institutional Interaction):""",

    23: """Empirical & Phenomenological Contribution: We document an observed pattern of AI-mediated information seeking in a public policy crisis context, examining how, when institutional accounts show limited interaction in a hyper-fragmented network (332 isolated clusters), users may seek information or interpretive assistance from AI-mediated accounts.""",

    117: """4.4 Structural Power Asymmetry: AI-Mediated Information Seeking vs. Institutional Interaction""",

    121: """| 1 | @grok | 0 | 42 | 0.000000 | AI-Mediated Broadcast Hub |""",

    157: """5.3 The Emergence of AI-Mediated Information Seeking""",

    166: """3. Conceptualization of AI-Mediated Information Seeking: We document an observed pattern in which an AI account occupies a central communicative position within the analyzed Indonesian public policy network, providing an empirical basis for examining the role of AI-mediated accounts in networked gatekeeping and information seeking."""
}

for n, replacement in replacements.items():
    idx = n - 1

    if idx >= len(doc.paragraphs):
        print(f"WARNING: P{n} does not exist")
        continue

    old = doc.paragraphs[idx].text

    if old != replacement:
        doc.paragraphs[idx].text = replacement
        changed.append((n, old, replacement))

doc.save(out)

print("HMC V21 CREATED")
print("=" * 100)
print("INPUT :", src)
print("OUTPUT:", out)
print("CHANGED PARAGRAPHS:", len(changed))
print()

for n, old, new in changed:
    print(f"P{n} CHANGED")
    print("OLD:", old)
    print("NEW:", new)
    print("-" * 100)
