from docx import Document
import re
import os

PATH = "HMC_REVIEW_MANUSCRIPT.docx"
doc = Document(PATH)
paras = doc.paragraphs
text = "\n".join(p.text for p in paras)

print("=" * 100)
print("HMC SUBMISSION QA — MASTER")
print("=" * 100)

print("\nFILE")
print("Path:", os.path.abspath(PATH))
print("Size:", round(os.path.getsize(PATH)/1024, 2), "KB")
print("Paragraphs:", len(paras))
print("Native Word tables:", len(doc.tables))
print("Sections:", len(doc.sections))

# ================================================================
# 1. ABSTRACT
# ================================================================
print("\n" + "=" * 100)
print("1. ABSTRACT")
print("=" * 100)

abstract_idx = None
for i, p in enumerate(paras):
    if p.text.strip().upper() == "ABSTRACT":
        abstract_idx = i
        break

if abstract_idx is not None:
    abstract_parts = []
    for p in paras[abstract_idx+1:]:
        t = p.text.strip()
        if t.upper().startswith("KEYWORDS"):
            break
        if t:
            abstract_parts.append(t)

    abstract = " ".join(abstract_parts)
    words = re.findall(r"\b[\w'-]+\b", abstract)

    print("Word count:", len(words))
    print("First 500 chars:")
    print(abstract[:500])

    if 150 <= len(words) <= 250:
        print("STATUS: PASS")
    else:
        print("STATUS: REVIEW — HMC target is 150–250 words")
else:
    print("STATUS: REVIEW — ABSTRACT heading not found")

# ================================================================
# 2. KEYWORDS
# ================================================================
print("\n" + "=" * 100)
print("2. KEYWORDS")
print("=" * 100)

for i, p in enumerate(paras):
    if p.text.strip().upper().startswith("KEYWORDS"):
        print(p.text.strip())
        break
else:
    print("REVIEW — Keywords not found")

# ================================================================
# 3. DOUBLE BLIND
# ================================================================
print("\n" + "=" * 100)
print("3. DOUBLE-BLIND")
print("=" * 100)

identity_terms = [
    "Indri Anjar Kartika Sari",
    "Universitas Pembangunan Nasional",
    "UPN Veteran Jawa Timur",
    "25067020011@student",
    "ORCID",
    "email:",
    "e-mail:"
]

hits = []
for i, p in enumerate(paras, 1):
    for term in identity_terms:
        if term.lower() in p.text.lower():
            hits.append((i, term))

if hits:
    for h in hits:
        print("REVIEW:", h)
else:
    print("PASS — no obvious author identity/contact text")

# ================================================================
# 4. LINE NUMBER / PAGE NUMBER
# ================================================================
print("\n" + "=" * 100)
print("4. REVIEW-MANUSCRIPT FEATURES")
print("=" * 100)

print("Note: python-docx cannot reliably determine visible Word line-number/page-number display.")
print("These must be checked visually in Word.")

# ================================================================
# 5. NATIVE TABLES VS PIPE TABLES
# ================================================================
print("\n" + "=" * 100)
print("5. TABLE CHECK")
print("=" * 100)

print("Native Word tables:", len(doc.tables))

pipe_rows = []
for i, p in enumerate(paras, 1):
    t = p.text.strip()
    if t.startswith("|") and t.endswith("|"):
        pipe_rows.append((i, t))

print("Pipe-format table paragraphs:", len(pipe_rows))

if pipe_rows:
    print("\nFirst pipe-table rows:")
    for i, t in pipe_rows[:15]:
        print(f"P{i}: {t[:250]}")

# ================================================================
# 6. FIGURE CALLOUTS
# ================================================================
print("\n" + "=" * 100)
print("6. FIGURE / TABLE CALLOUT SEARCH")
print("=" * 100)

for term in ["Figure", "Fig.", "Table", "TABLE", "FIGURE"]:
    count = len(re.findall(r"\b" + re.escape(term) + r"\b", text, flags=re.I))
    print(term, ":", count)

# ================================================================
# 7. REFERENCES
# ================================================================
print("\n" + "=" * 100)
print("7. REFERENCES")
print("=" * 100)

ref_idx = None
for i, p in enumerate(paras):
    if p.text.strip().upper() == "REFERENCES":
        ref_idx = i
        break

if ref_idx is None:
    print("REVIEW — REFERENCES heading not found")
else:
    refs = [p.text.strip() for p in paras[ref_idx+1:] if p.text.strip()]
    print("Reference paragraphs:", len(refs))

    print("\nFirst 10 references:")
    for r in refs[:10]:
        print("-", r[:300])

# ================================================================
# 8. APA / DOI / URL SIGNALS
# ================================================================
print("\n" + "=" * 100)
print("8. REFERENCE SIGNALS")
print("=" * 100)

if ref_idx is not None:
    ref_text = "\n".join(refs)

    print("DOI count :", len(re.findall(r"https?://doi\.org/", ref_text, re.I)))
    print("URL count :", len(re.findall(r"https?://", ref_text, re.I)))
    print("DOI without https://doi.org signal:",
          len(re.findall(r"\bdoi:\s*", ref_text, re.I)))

# ================================================================
# 9. SECTION / FORMATTING
# ================================================================
print("\n" + "=" * 100)
print("9. FORMATTING")
print("=" * 100)

for si, sec in enumerate(doc.sections, 1):
    print(f"Section {si}")
    print(" top   :", round(sec.top_margin.inches, 2))
    print(" bottom:", round(sec.bottom_margin.inches, 2))
    print(" left  :", round(sec.left_margin.inches, 2))
    print(" right :", round(sec.right_margin.inches, 2))

sample = [1,2,3,6,75,76,119,126,127,158,182,184]

for n in sample:
    if n <= len(paras):
        p = paras[n-1]
        run = p.runs[0] if p.runs else None

        print(
            f"P{n}: "
            f"style={p.style.name}; "
            f"font={run.font.name if run else None}; "
            f"size={run.font.size.pt if run and run.font.size else None}; "
            f"line_spacing={p.paragraph_format.line_spacing}; "
            f"before={p.paragraph_format.space_before}; "
            f"after={p.paragraph_format.space_after}"
        )

# ================================================================
# 10. CRITICAL TEXT CHECK
# ================================================================
print("\n" + "=" * 100)
print("10. CRITICAL TEXT")
print("=" * 100)

for n in [75,76,119,126,127,158,182,184]:
    if n <= len(paras):
        print(f"\nP{n}:")
        print(paras[n-1].text)

# ================================================================
# 11. HIGH-RISK WORDING
# ================================================================
print("\n" + "=" * 100)
print("11. HIGH-RISK WORDING")
print("=" * 100)

bad_terms = [
    "Algorithmic Oracle",
    "passive receptacles",
    "role abandoned",
    "most transformative",
    "dominating the graph",
    "remained silent",
    "absolute zero",
    "rancid physical meals",
    "severe trauma",
    "paralyze institutional legitimacy",
    "dismantle the conditions"
]

found = []

for i, p in enumerate(paras, 1):
    for term in bad_terms:
        if term.lower() in p.text.lower():
            found.append((i, term))

if found:
    for x in found:
        print("REVIEW:", x)
else:
    print("PASS — targeted high-risk expressions absent")

print("\n" + "=" * 100)
print("END HMC SUBMISSION QA")
print("=" * 100)
