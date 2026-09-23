from docx import Document
from zipfile import ZipFile
import re
import os

FILE = "HMC_REVIEW_MANUSCRIPT_V23.docx"
REPORT = "HMC_FINAL_QA_V24.txt"

doc = Document(FILE)
full_text = "\n".join(p.text for p in doc.paragraphs)

results = []
issues = []

def check(name, condition, detail=""):
    status = "PASS" if condition else "REVIEW"
    results.append((status, name, detail))
    if not condition:
        issues.append((name, detail))

# 1. FILE
check(
    "FILE EXISTS",
    os.path.exists(FILE),
    FILE
)

# 2. DOUBLE-BLIND
identity_patterns = [
    r"\bIndri Anjar Kartika Sari\b",
    r"\bUPN\b",
    r"Universitas Pembangunan Nasional",
    r"Veteran Jawa Timur",
    r"ORCID",
    r"25067020011@student",
    r"@student\.upnjatim",
]

identity_hits = []

for pattern in identity_patterns:
    if re.search(pattern, full_text, re.I):
        identity_hits.append(pattern)

check(
    "DOUBLE-BLIND IDENTITY",
    not identity_hits,
    "No author/affiliation/ORCID/contact identity found"
    if not identity_hits else str(identity_hits)
)

# 3. HIGH-RISK LANGUAGE
high_risk = [
    "Algorithmic Oracle",
    "rancid physical meals",
    "severe trauma",
    "paralyze institutional legitimacy",
    "dismantle the conditions",
    "passive receptacles",
    "absolute zero",
    "most transformative",
    "dominating the graph",
    "remained silent",
]

hits = []

for i, p in enumerate(doc.paragraphs, 1):
    for term in high_risk:
        if term.lower() in p.text.lower():
            hits.append(f"P{i}: {term}")

check(
    "HIGH-RISK LANGUAGE",
    not hits,
    "All targeted high-risk wording removed"
    if not hits else "\n".join(hits)
)

# 4. CONCATENATION
patterns = [
    r"\bandconsumer\b",
    r"\bpoliticalcommunication\b",
    r"\bwelfaremegaprojects\b",
    r"\bcommunicatedexpectations\b",
    r"\bGovernanceGap\b",
    r"\bSectorGovernance\b",
    r"\bandCode\b",
    r"\bofan\b",
    r"\binthe\b",
    r"\bofand\b",
    r"\btheand\b",
]

concat_hits = []

for i, p in enumerate(doc.paragraphs, 1):
    for pattern in patterns:
        if re.search(pattern, p.text, re.I):
            concat_hits.append(f"P{i}: {pattern}")

check(
    "SPACING / CONCATENATION",
    not concat_hits,
    "No targeted concatenation errors"
    if not concat_hits else "\n".join(concat_hits)
)

# 5. FONT
font_hits = []

for i, p in enumerate(doc.paragraphs, 1):
    for run in p.runs:
        if not run.text.strip():
            continue

        if run.font.name and run.font.name.lower() not in [
            "times new roman",
            "timesnewroman"
        ]:
            font_hits.append(f"P{i}: {run.font.name}")

check(
    "FONT",
    not font_hits,
    "Times New Roman"
    if not font_hits else "\n".join(font_hits[:20])
)

# 6. FONT SIZE
size_hits = []

for i, p in enumerate(doc.paragraphs, 1):
    for run in p.runs:
        if not run.text.strip():
            continue

        if run.font.size:
            size = run.font.size.pt
            if abs(size - 12) > 0.2:
                size_hits.append(f"P{i}: {size} pt")

check(
    "FONT SIZE",
    not size_hits,
    "12 pt"
    if not size_hits else "\n".join(size_hits[:20])
)

# 7. DOUBLE SPACING
line_hits = []

for i, p in enumerate(doc.paragraphs, 1):
    if not p.text.strip():
        continue

    spacing = p.paragraph_format.line_spacing

    if isinstance(spacing, float):
        if abs(spacing - 2.0) > 0.05:
            line_hits.append(
                f"P{i}: line_spacing={spacing}"
            )

check(
    "DOUBLE SPACING",
    not line_hits,
    "No explicit non-double spacing detected"
    if not line_hits else "\n".join(line_hits[:20])
)

# 8. MARGINS
margin_hits = []

for i, sec in enumerate(doc.sections, 1):
    values = {
        "top": sec.top_margin.inches,
        "bottom": sec.bottom_margin.inches,
        "left": sec.left_margin.inches,
        "right": sec.right_margin.inches,
    }

    for side, value in values.items():
        if abs(value - 1.0) > 0.05:
            margin_hits.append(
                f"Section {i} {side}={value:.3f}"
            )

check(
    "ONE-INCH MARGINS",
    not margin_hits,
    "All margins approximately 1 inch"
    if not margin_hits else "\n".join(margin_hits)
)

# 9. CORRECT MANUSCRIPT STRUCTURE
required_sections = [
    "ABSTRACT",
    "1. INTRODUCTION",
    "2. THEORETICAL FRAMEWORK AND LITERATURE REVIEW",
    "3. METHODOLOGY AND RESEARCH DESIGN",
    "4. EMPIRICAL RESULTS",
    "5. DISCUSSION",
    "6. THEORETICAL AND PRACTICAL CONTRIBUTIONS",
    "7. RESEARCH LIMITATIONS AND FUTURE RESEARCH AGENDA",
    "8. CONCLUSION",
    "REFERENCES",
]

missing_sections = []

for section in required_sections:
    if section.lower() not in full_text.lower():
        missing_sections.append(section)

check(
    "MANUSCRIPT STRUCTURE",
    not missing_sections,
    "All core manuscript sections detected"
    if not missing_sections
    else "Missing: " + ", ".join(missing_sections)
)

# 10. ALGORITHMIC ORACLE
oracle_count = len(
    re.findall(r"Algorithmic Oracle", full_text, re.I)
)

check(
    "ALGORITHMIC ORACLE REMOVED",
    oracle_count == 0,
    f"Occurrences: {oracle_count}"
)

# 11. CRITICAL RESULTS
critical_values = [
    "5,263",
    "3,395",
    "9.28%",
    "971",
    "666",
    "0.0011",
    "1.21%",
    "9.17%",
    "332",
    "0.9837",
    "56.24%",
    "20.39%",
    "12.33%",
    "9.60%",
]

missing_values = [
    value for value in critical_values
    if value not in full_text
]

check(
    "CRITICAL RESULTS PRESERVED",
    not missing_values,
    "Key reported values detected"
    if not missing_values
    else "Missing: " + ", ".join(missing_values)
)

# 12. TABLES
pipe_rows = []

for i, p in enumerate(doc.paragraphs, 1):
    txt = p.text.strip()

    if txt.startswith("|") and txt.endswith("|"):
        pipe_rows.append(i)

check(
    "TABLE FORMAT REVIEW",
    True,
    f"{len(pipe_rows)} pipe-style table rows detected; "
    "HMC table placement/format requires final visual review."
)

# 13. METADATA
metadata_hits = []

with ZipFile(FILE, "r") as z:
    core = z.read("docProps/core.xml").decode(
        "utf-8",
        errors="ignore"
    )

    for term in [
        "Indri",
        "Anjar",
        "Kartika",
        "UPN",
        "Veteran",
        "student.upnjatim",
    ]:
        if term.lower() in core.lower():
            metadata_hits.append(term)

check(
    "WORD METADATA ANONYMIZED",
    not metadata_hits,
    "No author identity detected in core metadata"
    if not metadata_hits else str(metadata_hits)
)

# ============================================================
# REPORT
# ============================================================

with open(REPORT, "w", encoding="utf-8") as f:
    f.write("HMC FINAL QA V24\n")
    f.write("=" * 100 + "\n")
    f.write(f"FILE: {FILE}\n\n")

    for status, name, detail in results:
        f.write(f"[{status}] {name}\n")
        f.write(f"       {detail}\n")

    f.write("\n" + "=" * 100 + "\n")

    if issues:
        f.write("FINAL STATUS: REVIEW REQUIRED\n")
        for name, detail in issues:
            f.write(f"- {name}: {detail}\n")
    else:
        f.write("FINAL STATUS: PASS\n")
        f.write("V23 PASSED ALL AUTOMATED HMC SUBMISSION CHECKS.\n")

print("=" * 100)
print("HMC FINAL QA V24")
print("=" * 100)

for status, name, detail in results:
    print(f"[{status}] {name}: {detail}")

print("\n" + "=" * 100)

if issues:
    print("FINAL STATUS: REVIEW REQUIRED")
    print("=" * 100)

    for name, detail in issues:
        print(f"- {name}")
        print(f"  {detail}")
else:
    print("FINAL STATUS: PASS")
    print("=" * 100)
    print("V23 PASSED ALL AUTOMATED HMC SUBMISSION CHECKS.")

print(f"\nREPORT: {REPORT}")
