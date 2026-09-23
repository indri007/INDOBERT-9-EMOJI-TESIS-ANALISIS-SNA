from docx import Document

src = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT.docx"
out = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V6.docx"

doc = Document(src)

replacements = {
    "communicationtopology": "communication topology",
    "orformalize": "or formalize",

    "$C_{      ext{out}} =42$": "$C_{\\text{out}} = 42$",
    "$C_{    ext{out}} = 0$": "$C_{\\text{out}} = 0$",

    "crisiscommunication": "crisis communication",
    "capability,while": "capability, while",

    "responsesto": "responses to",
    "TheAI": "The AI",
    "observedinteraction": "observed interaction",
    "anAI-mediated": "an AI-mediated",

    "Rather than establishing a historical paradigm shift, the finding provides an empirical instance of how algorithmic actors may participate in public-policy discourse alongside human institutional and non-institutional actors.":
    "The finding therefore provides an empirical instance of how algorithmic actors may participate in public-policy discourse alongside human institutional and non-institutional actors.",

    "asubstantial": "a substantial",
    "accountssuch": "accounts such",
}

changed = 0

for para in doc.paragraphs:
    original = para.text
    updated = original

    for old, new in replacements.items():
        updated = updated.replace(old, new)

    if updated != original:
        if para.runs:
            para.runs[0].text = updated
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.add_run(updated)

        changed += 1

print("CHANGED PARAGRAPHS:", changed)

doc.save(out)
print("OUTPUT:", out)
