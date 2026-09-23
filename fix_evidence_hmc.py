from docx import Document
from pathlib import Path

src = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT.docx")
out = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V3.docx")

doc = Document(src)

replacements = {
"""The network density of $0.0011$ and reciprocity of $1.21\\%$ indicate an absolute breakdown of conversational exchange. Citizens were not talking to one another, nor were institutions conversing with citizens. Instead, discourse manifested as outward broadcasting into an echo chamber void.""":

"""The network density of $0.0011$ and reciprocity of $1.21\\%$ indicate a substantial constraint on conversational exchange. Citizens were not frequently engaging in reciprocal interaction with one another, while institutional accounts showed limited reciprocal interaction with citizens in the observed network. Instead, discourse was characterized largely by outward broadcasting into localized clusters, resulting in a fragmented interaction structure.""",

"""Across all three operational aspects, Disgust exceeded $70\\%$, peaking at $78.91\\%$ in Logistics and $77.01\\%$ in Budget Allocation. This provides empirical proof of the Phygital Gap: negative affect was not an abstract political grievance, but an intense rejection of operational failures in delivery timing, budget misappropriation, and inadequate dietary portion sizes.""":

"""Across all three operational aspects, Disgust exceeded $70\\%$, peaking at $78.91\\%$ in Logistics and $77.01\\%$ in Budget Allocation. These findings provide empirical evidence consistent with the Phygital Gap interpretation: negative affect was associated with operational concerns expressed in the observed discourse, including delivery timing, budget allocation, and dietary portion sizes.""",

"""The empirical findings substantiate the central theoretical proposition of this study: the crisis of the MBG program represents a textbook manifestation of the Phygital Gap (Kotler et al., 2023) transposed into public policy communication. In corporate marketing, a phygital gap erodes brand equity; in public governance, it destroys democratic legitimacy. The government succeeded in crafting an appealing digital vision of national nutritional sovereignty across press releases, social media infographics, and broadcast announcements. However, when the physical deliverable—the meal tray—arrived at schools in unhygienic states or with minimal portions, citizens experienced sharp sensory betrayal.""":

"""The empirical findings provide evidence relevant to the central theoretical proposition of this study: the crisis of the MBG program can be interpreted through the Phygital Gap (Kotler et al., 2023) in the context of public policy communication. In corporate marketing, a phygital gap concerns discrepancies between digital representations and physical or experiential delivery; in public governance, similar discrepancies may affect perceptions of policy credibility and institutional legitimacy. The government communicated an appealing digital vision of national nutritional sovereignty through press releases, social media infographics, and broadcast announcements. However, when the physical deliverable—the meal tray—was perceived by users as unhygienic or insufficient in portion size, the discrepancy between communicated expectations and reported physical experiences became salient in public discourse.""",

"""The dominance of Disgust (56.24%) over Anger (1.05%) is sociologically profound. In political psychology, anger is an approach-oriented emotion that motivates political activism, protest, and demands for institutional reform. Disgust, conversely, is an avoidance-oriented visceral emotion triggered by physical contaminants, spoiled food, and moral revulsion. By failing at the physical touchpoint, the government reduced a multi-trillion-rupiah national policy into an object of physical disgust.""":

"""The dominance of Disgust (56.24%) over Anger (1.05%) is analytically significant within the observed corpus. In political psychology, anger is generally associated with approach-oriented responses that can motivate political activism, protest, and demands for institutional reform, whereas disgust is associated with avoidance-oriented responses and may be triggered by perceptions of contamination or moral violation. In the context of the present study, the predominance of Disgust may indicate that negative public responses were more strongly associated with perceived problems at the physical and operational touchpoints of the MBG program than with mobilization-oriented anger. This pattern should be interpreted as a characteristic of the observed corpus rather than as evidence of a general psychological response to the policy.""",

"""Coupled with a reciprocity rate of only 1.21%, this reveals that the digital discourse was functionally dead as a deliberative space. Citizens broadcasted their frustration into localized micro-clusters, while government actors failed to engage in reciprocal dialogue. The giant component captured only $9.17\\%$ of nodes, indicating that no single narrative could bridge the fragmented silos.""":

"""Coupled with a reciprocity rate of only 1.21%, this suggests that the observed digital discourse was structurally constrained as a deliberative space. Citizens broadcasted their frustration into localized micro-clusters, while reciprocal interaction with government actors was limited in the observed network. The giant component captured only $9.17\\%$ of nodes, indicating that the network contained substantial fragmentation and limited connectivity across discourse clusters.""",

"""When state actors create a Power Vacuum by failing to respond to legitimate operationalqueries, digital citizens do not abandon inquiry; instead, they pivot to synthetic authority. Netizens tagged @grok to verify contract values, calculate vendor margins, and cross-reference nutritional standards. The AI agent operated as an impartial Algorithmic Oracle, synthesizing facts in real time. Thissignals a historic paradigm shift in digital public relations: future state communication strategies will not merely interact with human journalists or influencers, but must contend with autonomous algorithms operating as primary epistemic gatekeepers in the network.""":

"""When state actors provide limited responses to operational queries, some users may redirect information-seeking activities toward AI-mediated accounts. In the observed network, users tagged @grok in attempts to verify contract values, calculate vendor margins, and cross-reference nutritional standards. The AI agent therefore occupied a distinct information-seeking position within the observed interaction network. This pattern is relevant to human-machine communication because an AI-mediated account can become part of the communicative infrastructure through which users seek, interpret, and validate policy-related information. Rather than establishing a historical paradigm shift, the finding provides an empirical instance of how algorithmic actors may participate in public-policy discourse alongside human institutional and non-institutional actors.""",

"""1. Closing the Physical Touchpoint: Public sector communication teams cannot cure an operational crisis with digital spin. PR resources must be directed toward auditing ground-level operational logistics before launching promotional digital campaigns.""":

"""1. Closing the Physical Touchpoint: Public sector communication teams are unlikely to resolve an operational crisis through digital communication alone. PR resources should therefore be complemented by attention to ground-level operational logistics before launching promotional digital campaigns.""",

"""2. Eliminating the Power Vacuum: Institutional communicators must abandon one-way broadcast practices. Establishing responsive, interactive social listening units capable of engaging with citizen inquiries is essential to raise network reciprocity and restore institutional trust.""":

"""2. Reducing the Power Vacuum: Institutional communicators may need to reconsider one-way broadcast practices in contexts characterized by limited reciprocal interaction. Establishing responsive, interactive social listening units capable of engaging with citizen inquiries could help increase network reciprocity and support institutional trust.""",

"""3. Deprecating Binary Sentiment Tools: Government monitoring dashboards must replace obsolete positive/negative lexicon tools with contextual transformer models capable of decoding sarcasm and granular disgust.""":

"""3. Moving Beyond Binary Sentiment Tools: Government monitoring dashboards may benefit from complementing positive/negative lexicon approaches with contextual transformer models capable of identifying sarcasm and more granular affective categories, including disgust."""
}

changed = 0

for para in doc.paragraphs:
    text = para.text
    if text in replacements:
        # preserve paragraph formatting while replacing runs
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = replacements[text]
        else:
            para.add_run(replacements[text])
        changed += 1

doc.save(out)

print(f"CHANGED PARAGRAPHS: {changed}")
print(f"OUTPUT: {out}")
