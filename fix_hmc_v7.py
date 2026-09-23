from docx import Document
import re

src = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V6.docx"
out = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V7.docx"

doc = Document(src)

changed = 0

replacements = {
    "communicationtopology": "communication topology",
    "orformalize": "or formalize",
    "Government figuresacted": "Government figures acted",
    "crisiscommunication": "crisis communication",
    "capability,while": "capability, while",
    "responsesto": "responses to",
    "TheAI": "The AI",
    "observedinteraction": "observed interaction",
    "anAI-mediated": "an AI-mediated",
    "asubstantial": "a substantial",
    "accountssuch": "accounts such",
}

for i, para in enumerate(doc.paragraphs, 1):

    original = para.text
    updated = original

    # Confirmed spacing corrections
    for old, new in replacements.items():
        updated = updated.replace(old, new)

    # Repair malformed C_out
    updated = re.sub(
        r'\$C_\{\s*ext\{out\}\}\s*=\s*([0-9]+)\$',
        r'$C_{\\text{out}} = \1$',
        updated
    )

    # Repair malformed C_out(v)
    updated = re.sub(
        r'\$C_\{\s*ext\{out\}\}\(v\)\$',
        r'$C_{\\text{out}}(v)$',
        updated
    )

    # Repair malformed C_in
    updated = re.sub(
        r'\$C_\{\s*ext\{in\}\}\$',
        r'$C_{\\text{in}}$',
        updated
    )

    if updated != original:

        if para.runs:
            para.runs[0].text = updated

            for run in para.runs[1:]:
                run.text = ""
        else:
            para.add_run(updated)

        changed += 1
        print(f"CHANGED P{i}")

doc.save(out)

print()
print("CHANGED PARAGRAPHS:", changed)
print("OUTPUT:", out)
