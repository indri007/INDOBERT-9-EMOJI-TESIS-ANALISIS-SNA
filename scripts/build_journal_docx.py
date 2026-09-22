import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_formatted_text(p, text):
    # Regex to handle bold, italic, code, and plain text
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(40, 40, 120)
        else:
            p.add_run(token)

def build_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = docx.Document()

    # Set 1 inch margins (APA standard)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    lines = content.splitlines()
    i = 0
    total_lines = len(lines)

    while i < total_lines:
        line = lines[i].strip()

        # Skip empty lines or horizontal rules
        if not line or line == '---':
            i += 1
            continue

        # Title (# )
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(title_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(15, 30, 60)
            i += 1
            continue

        # Heading 1 (## )
        if line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(20, 40, 80)
            i += 1
            continue

        # Heading 2 (### )
        if line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(30, 50, 90)
            i += 1
            continue

        # Heading 3 (#### )
        if line.startswith('#### '):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(h_text)
            run.bold = True
            run.italic = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            i += 1
            continue

        # Image ![caption](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            if img_path.startswith('file://'):
                img_path = img_path[7:]
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.add_run().add_picture(img_path, width=Inches(5.8))
                
                # Add Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(12)
                cap_run = p_cap.add_run(caption)
                cap_run.italic = True
                cap_run.font.size = Pt(10)
                cap_run.font.color.rgb = RGBColor(60, 60, 60)
            i += 1
            continue

        # Blockquote (> )
        if line.startswith('> '):
            quote_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_formatted_text(p, quote_text)
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(40, 40, 40)
            i += 1
            continue

        # Code block (```)
        if line.startswith('```'):
            i += 1
            code_lines = []
            while i < total_lines and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1 # skip closing ```
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            code_run = p.add_run("\n".join(code_lines))
            code_run.font.name = 'Courier New'
            code_run.font.size = Pt(9.5)
            code_run.font.color.rgb = RGBColor(30, 30, 30)
            continue

        # Table detection (| col1 | col2 |)
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            # Parse table
            rows_data = []
            for tl in table_lines:
                # Check if it's separator row |:---|:---:|
                if re.match(r'^\|[\s:\-\|]+$', tl):
                    continue
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                rows_data.append(cells)

            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True

                for r_idx, r_data in enumerate(rows_data):
                    row = tbl.rows[r_idx]
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = row.cells[c_idx]
                        val = r_data[c_idx] if c_idx < len(r_data) else ""
                        val = val.replace('<br>', '\n')
                        cell.text = ""
                        p_cell = cell.paragraphs[0]
                        p_cell.paragraph_format.space_before = Pt(3)
                        p_cell.paragraph_format.space_after = Pt(3)
                        p_cell.paragraph_format.line_spacing = 1.05
                        add_formatted_text(p_cell, val)

                        if is_header:
                            set_cell_background(cell, "EAEFF5")
                            for r in p_cell.runs:
                                r.bold = True
                                r.font.size = Pt(10)
                                r.font.color.rgb = RGBColor(10, 30, 60)
                        else:
                            for r in p_cell.runs:
                                r.font.size = Pt(9.5)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Bullet list (- )
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, line[2:].strip())
            i += 1
            continue

        # Numbered list (1. , 2. )
        num_match = re.match(r'^([0-9]+)\.\s+(.*)$', line)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, num_match.group(2).strip())
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        # Author / Affiliation styling
        if line.startswith('**Indri Anjar Kartika Sari**') or 'UPN Veteran' in line or 'Email:' in line:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
        add_formatted_text(p, line)
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated formal manuscript: {docx_path}")

if __name__ == '__main__':
    md_file = '/Users/jevin/Documents/tesis_mbg/journal_paper_mbg_sna.md'
    docx_file = '/Users/jevin/Documents/tesis_mbg/Journal_Paper_Indri_Anjar_MBG_SNA.docx'
    build_docx(md_file, docx_file)
