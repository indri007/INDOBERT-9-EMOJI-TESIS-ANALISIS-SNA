from pathlib import Path
from copy import deepcopy
import re

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# CONFIGURATION
# ============================================================

CANDIDATES = [
    Path("/Users/jevin/jurnal/output/MANUSCRIPT_SCOPUS_Q1_FINAL.docx"),
    Path("/Users/jevin/jurnal/output/MANUSCRIPT_FINAL_REVISED.docx"),
    Path("/Users/jevin/jurnal/output/MANUSCRIPT_SOURCE.docx"),
    Path("/Users/jevin/Documents/tesis_mbg/manuscript/Anonymized_Manuscript_Scopus_Q1.docx"),
]

OUTPUT_DIR = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission")

OUTPUT_DOCX = OUTPUT_DIR / "HMC_REVIEW_MANUSCRIPT.docx"
AUDIT_FILE = OUTPUT_DIR / "HMC_MANUSCRIPT_AUDIT.txt"


# ============================================================
# FIND SOURCE
# ============================================================

def find_source():
    for path in CANDIDATES:
        if path.exists():
            return path

    raise FileNotFoundError(
        "Tidak ditemukan manuscript DOCX.\n"
        "Periksa path pada CANDIDATES."
    )


# ============================================================
# WORD METADATA
# ============================================================

def clear_metadata(document):
    props = document.core_properties

    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.subject = ""

    # Jangan menghapus title manuscript.
    # Title tetap berada di body dokumen.


# ============================================================
# PAGE NUMBERS
# ============================================================

def add_page_number(paragraph):
    run = paragraph.add_run()

    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "

    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


# ============================================================
# CONTINUOUS LINE NUMBERS
# ============================================================

def add_continuous_line_numbers(document):
    sect_pr = document.sections[0]._sectPr

    existing = sect_pr.find(qn("w:lnNumType"))

    if existing is not None:
        sect_pr.remove(existing)

    ln_num = OxmlElement("w:lnNumType")

    # continuous numbering
    ln_num.set(qn("w:countBy"), "1")
    ln_num.set(qn("w:start"), "1")
    ln_num.set(qn("w:distance"), "360")
    ln_num.set(qn("w:restart"), "continuous")

    sect_pr.append(ln_num)


# ============================================================
# PAGE SETUP
# ============================================================

def apply_page_setup(document):

    for section in document.sections:

        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # PAGE NUMBER
        footer = section.footer

        if not footer.paragraphs:
            paragraph = footer.add_paragraph()
        else:
            paragraph = footer.paragraphs[0]

        paragraph.alignment = 1

        # Avoid duplicate page fields
        paragraph.clear()

        add_page_number(paragraph)


# ============================================================
# FONT + DOUBLE SPACING
# ============================================================

def apply_hmc_formatting(document):

    for paragraph in document.paragraphs:

        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        for run in paragraph.runs:

            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

            # Ensure East Asian font also uses Times New Roman
            rpr = run._r.get_or_add_rPr()

            rfonts = rpr.rFonts

            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)

            rfonts.set(qn("w:ascii"), "Times New Roman")
            rfonts.set(qn("w:hAnsi"), "Times New Roman")
            rfonts.set(qn("w:eastAsia"), "Times New Roman")


# ============================================================
# MARKDOWN CLEANUP
# ============================================================

def clean_markdown_formatting(document):

    patterns = [
        (r"\*\*(.*?)\*\*", r"\1"),
        (r"__(.*?)__", r"\1"),
        (r"`(.*?)`", r"\1"),
    ]

    for paragraph in document.paragraphs:

        for run in paragraph.runs:

            text = run.text

            for pattern, replacement in patterns:
                text = re.sub(
                    pattern,
                    replacement,
                    text
                )

            run.text = text


# ============================================================
# IDENTIFYING INFORMATION
# ============================================================

IDENTIFYING_TERMS = [
    "Indri Anjar Kartika Sari",
    "UPN Veteran Jawa Timur",
    "Universitas Pembangunan Nasional",
    "Universitas Pembangunan Nasional \"Veteran\" Jawa Timur",
    "Universitas Pembangunan Nasional 'Veteran' Jawa Timur",
]


def audit_identifying_information(document):

    findings = []

    for i, paragraph in enumerate(document.paragraphs, start=1):

        text = paragraph.text

        for term in IDENTIFYING_TERMS:

            if term.lower() in text.lower():

                findings.append(
                    f"Paragraph {i}: {term}"
                )

    return findings


# ============================================================
# ANONYMIZE IDENTIFYING INFORMATION
# ============================================================

def anonymize(document):

    replacements = {
        "Indri Anjar Kartika Sari": "[AUTHOR NAME REMOVED]",
        "Universitas Pembangunan Nasional \"Veteran\" Jawa Timur":
            "[AFFILIATION REMOVED]",
        "Universitas Pembangunan Nasional 'Veteran' Jawa Timur":
            "[AFFILIATION REMOVED]",
        "Universitas Pembangunan Nasional":
            "[AFFILIATION REMOVED]",
        "UPN Veteran Jawa Timur":
            "[AFFILIATION REMOVED]",
    }

    for paragraph in document.paragraphs:

        for run in paragraph.runs:

            text = run.text

            for old, new in replacements.items():

                text = text.replace(
                    old,
                    new
                )

            run.text = text


# ============================================================
# ABSTRACT
# ============================================================

def get_abstract_word_count(document):

    paragraphs = document.paragraphs

    start = None
    end = None

    for i, paragraph in enumerate(paragraphs):

        text = paragraph.text.strip().upper()

        if text == "ABSTRACT":
            start = i + 1
            continue

        if start is not None and re.match(
            r"^(KEYWORDS?|1\.|INTRODUCTION)",
            text
        ):
            end = i
            break

    if start is None:
        return None

    if end is None:
        end = len(paragraphs)

    abstract = " ".join(
        p.text for p in paragraphs[start:end]
    )

    return len(
        re.findall(r"\b[\w'-]+\b", abstract)
    )


# ============================================================
# KEYWORDS
# ============================================================

def get_keyword_count(document):

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if re.match(
            r"^(KEYWORDS?|Keywords?)\s*:",
            text,
            re.IGNORECASE
        ):

            value = re.sub(
                r"^(KEYWORDS?|Keywords?)\s*:",
                "",
                text,
                flags=re.IGNORECASE
            )

            keywords = [
                x.strip()
                for x in re.split(r"[;,]", value)
                if x.strip()
            ]

            return len(keywords)

    return None


# ============================================================
# UNDERLINE AUDIT
# ============================================================

def audit_underlines(document):

    findings = []

    for i, paragraph in enumerate(
        document.paragraphs,
        start=1
    ):

        for run in paragraph.runs:

            if run.underline:

                findings.append(
                    f"Paragraph {i}: {run.text}"
                )

    return findings


# ============================================================
# STRONG CLAIM AUDIT
# ============================================================

STRONG_TERMS = [
    "absolute breakdown",
    "empirical proof",
    "functionally dead",
    "historic paradigm shift",
    "destroys democratic legitimacy",
    "sociologically profound",
    "textbook manifestation",
    "must abandon",
    "obsolete",
    "cannot cure",
    "definitive proof",
    "proves",
]

def audit_strong_claims(document):

    findings = []

    for i, paragraph in enumerate(
        document.paragraphs,
        start=1
    ):

        text = paragraph.text.lower()

        for term in STRONG_TERMS:

            if term.lower() in text:

                findings.append(
                    f"Paragraph {i}: {term}"
                )

    return findings


# ============================================================
# SECTION AUDIT
# ============================================================

REQUIRED_SECTIONS = [
    "ABSTRACT",
    "INTRODUCTION",
    "METHODOLOGY",
    "RESULTS",
    "DISCUSSION",
    "CONCLUSION",
    "REFERENCES",
]

def audit_sections(document):

    text = "\n".join(
        p.text.upper()
        for p in document.paragraphs
    )

    result = {}

    for section in REQUIRED_SECTIONS:

        result[section] = section in text

    return result


# ============================================================
# REFERENCES AUDIT
# ============================================================

def audit_references(document):

    text = "\n".join(
        p.text
        for p in document.paragraphs
    )

    references_found = "REFERENCES" in text.upper()

    return references_found


# ============================================================
# SAVE AUDIT
# ============================================================

def save_audit(
    source,
    abstract_count,
    keyword_count,
    identifying,
    underlines,
    strong_claims,
    sections,
    references,
):

    lines = []

    lines.append("=" * 70)
    lines.append("HMC MANUSCRIPT SUBMISSION AUDIT")
    lines.append("=" * 70)
    lines.append("")

    lines.append(f"SOURCE:")
    lines.append(str(source))
    lines.append("")

    lines.append("OUTPUT:")
    lines.append(str(OUTPUT_DOCX))
    lines.append("")

    lines.append("=" * 70)
    lines.append("FORMATTING")
    lines.append("=" * 70)

    lines.append("Font: Times New Roman 12 pt")
    lines.append("Line spacing: DOUBLE")
    lines.append("Margins: 1 inch")
    lines.append("Page numbers: ENABLED")
    lines.append("Continuous line numbers: ENABLED")
    lines.append("")

    lines.append("=" * 70)
    lines.append("ABSTRACT")
    lines.append("=" * 70)

    if abstract_count is None:
        lines.append("STATUS: NOT DETECTED")
    else:
        lines.append(
            f"Word count: {abstract_count}"
        )

        if 150 <= abstract_count <= 250:
            lines.append("STATUS: PASS")
        else:
            lines.append("STATUS: NEEDS REVISION")

    lines.append("")

    lines.append("=" * 70)
    lines.append("KEYWORDS")
    lines.append("=" * 70)

    if keyword_count is None:
        lines.append("STATUS: NOT DETECTED")
    else:
        lines.append(
            f"Keyword count: {keyword_count}"
        )

        if 4 <= keyword_count <= 6:
            lines.append("STATUS: PASS")
        else:
            lines.append("STATUS: NEEDS REVISION")

    lines.append("")

    lines.append("=" * 70)
    lines.append("ANONYMIZATION")
    lines.append("=" * 70)

    if identifying:
        lines.append("IDENTIFYING TERMS FOUND BEFORE ANONYMIZATION:")
        for item in identifying:
            lines.append(item)
    else:
        lines.append(
            "No identifying terms detected."
        )

    lines.append("")

    lines.append("=" * 70)
    lines.append("UNDERLINE AUDIT")
    lines.append("=" * 70)

    if underlines:
        for item in underlines:
            lines.append(item)
        lines.append("STATUS: NEEDS REVIEW")
    else:
        lines.append(
            "No underlined text detected."
        )
        lines.append("STATUS: PASS")

    lines.append("")

    lines.append("=" * 70)
    lines.append("SECTION AUDIT")
    lines.append("=" * 70)

    for section, status in sections.items():

        lines.append(
            f"{section}: "
            f"{'PASS' if status else 'NEEDS REVIEW'}"
        )

    lines.append("")

    lines.append("=" * 70)
    lines.append("REFERENCES")
    lines.append("=" * 70)

    lines.append(
        "References section detected: "
        + ("YES" if references else "NO")
    )

    lines.append(
        "Full APA 7 bibliographic verification requires manual/source-level checking."
    )

    lines.append("")

    lines.append("=" * 70)
    lines.append("STRONG CLAIM AUDIT")
    lines.append("=" * 70)

    if strong_claims:

        for item in strong_claims:
            lines.append(item)

        lines.append("")
        lines.append(
            "STATUS: NEEDS AUTHOR REVIEW"
        )

    else:

        lines.append(
            "No predefined strong-claim terms detected."
        )

    lines.append("")

    lines.append("=" * 70)
    lines.append("REQUIRED HMC STATEMENTS")
    lines.append("=" * 70)

    lines.append(
        "Funding: AUTHOR VERIFICATION REQUIRED"
    )

    lines.append(
        "Conflict of Interest: AUTHOR VERIFICATION REQUIRED"
    )

    lines.append(
        "Data Availability: AUTHOR VERIFICATION REQUIRED"
    )

    lines.append(
        "Ethics: AUTHOR VERIFICATION REQUIRED"
    )

    lines.append(
        "CRediT: AUTHOR VERIFICATION REQUIRED"
    )

    lines.append(
        "AI / Tool Disclosure: AUTHOR VERIFICATION REQUIRED"
    )

    lines.append(
        "Acknowledgments: AUTHOR VERIFICATION REQUIRED"
    )

    lines.append("")

    lines.append("=" * 70)
    lines.append("IMPORTANT")
    lines.append("=" * 70)

    lines.append(
        "This audit does not invent missing bibliographic, ethics,"
    )

    lines.append(
        "funding, authorship, or data-availability information."
    )

    AUDIT_FILE.write_text(
        "\n".join(lines),
        encoding="utf-8"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print("HMC MANUSCRIPT FORMATTER")
    print("=" * 70)

    source = find_source()

    print("\nSOURCE:")
    print(source)

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    document = Document(
        str(source)
    )

    # Audit BEFORE anonymization
    identifying = audit_identifying_information(
        document
    )

    underlines = audit_underlines(
        document
    )

    strong_claims = audit_strong_claims(
        document
    )

    # Apply formatting
    clear_metadata(document)

    anonymize(document)

    clean_markdown_formatting(document)

    apply_hmc_formatting(document)

    apply_page_setup(document)

    add_continuous_line_numbers(document)

    # Audit content
    abstract_count = get_abstract_word_count(
        document
    )

    keyword_count = get_keyword_count(
        document
    )

    sections = audit_sections(
        document
    )

    references = audit_references(
        document
    )

    # Save
    document.save(
        str(OUTPUT_DOCX)
    )

    save_audit(
        source=source,
        abstract_count=abstract_count,
        keyword_count=keyword_count,
        identifying=identifying,
        underlines=underlines,
        strong_claims=strong_claims,
        sections=sections,
        references=references,
    )

    print("\n" + "=" * 70)
    print("SELESAI")
    print("=" * 70)

    print("\nOUTPUT FILE:")
    print(OUTPUT_DOCX)

    print("\nAUDIT FILE:")
    print(AUDIT_FILE)

    print("\nABSTRACT:")
    print(abstract_count, "words")

    print("\nKEYWORDS:")
    print(keyword_count)

    print("\nIDENTIFYING TERMS BEFORE CLEANUP:")
    print(len(identifying))

    print("\nSTRONG CLAIMS:")
    print(len(strong_claims))

    print("\nHMC formatter selesai.")


if __name__ == "__main__":
    main()
