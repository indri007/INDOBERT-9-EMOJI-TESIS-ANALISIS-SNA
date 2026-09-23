from docx import Document
import re
from pathlib import Path

p = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V15.docx")
doc = Document(p)

print("=" * 100)
print("FINAL HMC AUDIT — V15")
print("=" * 100)

print("\nFILE")
print("Path:", p)
print("Size:", round(p.stat().st_size / 1024, 1), "KB")

print("\nDOCUMENT STRUCTURE")
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
print("Sections:", len(doc.sections))

# ---------------------------------------------------------
# 1. DOUBLE BLIND / IDENTITY
# ---------------------------------------------------------
text = "\n".join(x.text for x in doc.paragraphs)

identity_patterns = [
    r'Indri Anjar Kartika Sari',
    r'UPN Veteran Jawa Timur',
    r'Universitas Pembangunan Nasional',
    r'25067020011',
    r'@student\.upnjatim\.ac\.id',
    r'ORCID',
    r'Authors:',
    r'Affiliation:',
]

print("\n" + "=" * 100)
print("1. DOUBLE-BLIND IDENTITY CHECK")
print("=" * 100)

hits = []
for pat in identity_patterns:
    if re.search(pat, text, re.I):
        hits.append(pat)

if hits:
    print("REVIEW:", hits)
else:
    print("PASS — no obvious author identity markers found.")

# ---------------------------------------------------------
# 2. INTERNAL METADATA
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("2. INTERNAL METADATA CHECK")
print("=" * 100)

metadata = [
    "Target Publication:",
    "Scopus ASJC Classification:",
    "Scopus Subject Area:",
    "Social Sciences (Communication)",
    "Computer Science (Artificial Intelligence)",
]

meta_hits = [x for x in metadata if x.lower() in text.lower()]

if meta_hits:
    print("REVIEW:", meta_hits)
else:
    print("PASS — internal publication metadata absent.")

# ---------------------------------------------------------
# 3. DECLARATIONS
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("3. REQUIRED DECLARATIONS")
print("=" * 100)

declarations = [
    "Funding",
    "Competing Interests",
    "Conflict of Interest",
    "Data and Code Availability",
    "Author Contributions",
    "CRediT",
    "Ethics Approval",
    "AI",
]

for d in declarations:
    found = d.lower() in text.lower()
    print(f"{d:30}:", "PASS" if found else "REVIEW")

# ---------------------------------------------------------
# 4. MALFORMED LATEX
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("4. LATEX / MATH CORRUPTION CHECK")
print("=" * 100)

bad_patterns = [
    r'ext\{out\}',
    r'ext\{in\}',
    r'C_\{\s*ext',
    r'\\text\s*\{\s*\}',
    r'\{\s+ext\{',
]

latex_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    for pat in bad_patterns:
        if re.search(pat, para.text):
            latex_hits.append((i, pat, para.text))

if latex_hits:
    for i, pat, line in latex_hits:
        print(f"P{i} | {pat}")
        print(line)
else:
    print("PASS — no known malformed LaTeX patterns.")

# ---------------------------------------------------------
# 5. TARGETED SPACING
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("5. TARGETED SPACING CHECK")
print("=" * 100)

spacing_patterns = [
    r'\b\w+are\b',
    r'\b\w+over\b',
    r'\bprimaryarena\b',
    r'\bdarkirony\b',
    r'\bpositive\'or\b',
    r'\bcommunicationtopology\b',
    r'\bwithpublic\b',
    r'\bofnetworked\b',
    r'\bnetworkedgatekeeping\b',
    r'\blaunchingpromotional\b',
    r'\bpromotionaldigital\b',
    r'\bpracticesin\b',
    r'\bfromcomplementing\b',
    r'\bfinalcorpus\b',
    r'\btheanalyzed\b',
    r'\bFAIRdata\b',
    r'\bseparateTitle\b',
    r'\bprovidedin\b',
    r'\bforDouble\b',
    r'\bCorporateReputation\b',
    r'\bofdeep\b',
    r'\binquiryinto\b',
    r'\bofSciences\b',
    r'\bIn2018\b',
    r'MBG\)program',
    r'103\(23\),8577',
]

spacing_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    for pat in spacing_patterns:
        if re.search(pat, para.text, re.I):
            spacing_hits.append((i, pat, para.text))

if spacing_hits:
    for i, pat, line in spacing_hits:
        print(f"P{i} | {pat}")
        print(line)
else:
    print("PASS — no known targeted spacing errors.")

# ---------------------------------------------------------
# 6. STRONG CLAIM REVIEW
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("6. STRONG CLAIM REVIEW")
print("=" * 100)

strong_patterns = [
    r'\bfundamentally altered\b',
    r'\binevitably\b',
    r'\bseverely misleading\b',
    r'\bcompels\b',
    r'\babandon human authorities\b',
    r'\bstate truth\b',
]

strong_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    for pat in strong_patterns:
        if re.search(pat, para.text, re.I):
            strong_hits.append((i, pat, para.text))

if strong_hits:
    for i, pat, line in strong_hits:
        print(f"P{i} | {pat}")
        print(line)
else:
    print("PASS — previous high-risk wording not found.")

# ---------------------------------------------------------
# 7. HMC STRUCTURAL HEADINGS
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("7. MAIN STRUCTURE")
print("=" * 100)

required_sections = [
    "INTRODUCTION",
    "METHOD",
    "RESULTS",
    "DISCUSSION",
    "LIMITATIONS",
    "CONCLUSION",
    "REFERENCES",
]

for s in required_sections:
    found = s.lower() in text.lower()
    print(f"{s:20}:", "PASS" if found else "REVIEW")

# ---------------------------------------------------------
# 8. FORMATTING
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("8. FORMATTING SAMPLE")
print("=" * 100)

for i in [0, 1, 5, 7]:
    if i < len(doc.paragraphs):
        p0 = doc.paragraphs[i]
        fmt = p0.paragraph_format
        run = p0.runs[0] if p0.runs else None

        print(f"P{i+1}")
        print("  style:", p0.style.name)
        print("  line spacing:", fmt.line_spacing)
        print("  before:", fmt.space_before)
        print("  after:", fmt.space_after)

        if run:
            print("  font:", run.font.name)
            print("  size:", run.font.size)

# margins
for i, sec in enumerate(doc.sections, 1):
    print(f"\nSection {i} margins:")
    print(" top   :", sec.top_margin.inches)
    print(" bottom:", sec.bottom_margin.inches)
    print(" left  :", sec.left_margin.inches)
    print(" right :", sec.right_margin.inches)

# ---------------------------------------------------------
# 9. KEY PARAGRAPHS
# ---------------------------------------------------------
print("\n" + "=" * 100)
print("9. KEY PARAGRAPH VERIFICATION")
print("=" * 100)

for n in [6, 8, 18, 119]:
    if n <= len(doc.paragraphs):
        print(f"\nP{n}:")
        print(doc.paragraphs[n-1].text)

print("\n" + "=" * 100)
print("FINAL AUDIT COMPLETE")
print("=" * 100)
