from docx import Document
from pathlib import Path

src = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V8.docx")
out = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V9.docx")

doc = Document(str(src))

replacements = {
    "communicationtopology": "communication topology",
    "orformalize": "or formalize",

    "$C_{      ext{in}}(v)$": "$C_{\\text{in}}(v)$",
    "$C_{      ext{out}}$": "$C_{\\text{out}}$",
    "$C_{      ext{out}}(v)$": "$C_{\\text{out}}(v)$",

    "facts.The": "facts. The",
    "crisiscommunication": "crisis communication",
    "capability,while": "capability, while",
}

changed = 0

for i, para in enumerate(doc.paragraphs, 1):
    original = para.text
    new_text = original

    for old, new in replacements.items():
        new_text = new_text.replace(old, new)

    if new_text != original:
        # Preserve paragraph formatting as much as possible.
        # Rebuild text in the paragraph while retaining paragraph-level formatting.
        for run in para.runs:
            run.text = ""

        if para.runs:
            para.runs[0].text = new_text
        else:
            para.add_run(new_text)

        changed += 1
        print(f"CHANGED P{i}")

doc.save(str(out))

print()
print(f"CHANGED PARAGRAPHS: {changed}")
print(f"OUTPUT: {out}")
