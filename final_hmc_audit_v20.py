from docx import Document
import re
import os

p = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V20.docx"
doc = Document(p)

print("=" * 100)
print("FINAL HMC AUDIT V20")
print("=" * 100)

print("\nFILE")
print("Path:", p)
print("Size:", round(os.path.getsize(p)/1024, 2), "KB")

print("\nDOCUMENT")
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
print("Sections:", len(doc.sections))

# ------------------------------------------------------------------
# 1. DOUBLE-BLIND / METADATA
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("1. DOUBLE-BLIND / INTERNAL METADATA")
print("=" * 100)

for i, para in enumerate(doc.paragraphs, 1):
    t = para.text.strip()

    forbidden = [
        "Indri Anjar Kartika Sari",
        "Universitas Pembangunan Nasional",
        "UPN Veteran",
        "25067020011",
        "Authors:",
        "Target Publication:",
        "Scopus ASJC Classification:",
        "Scopus Subject Area:"
    ]

    hits = [x for x in forbidden if x.lower() in t.lower()]

    if hits:
        print(f"REVIEW P{i}: {hits}")

print("Double-blind textual identity check complete.")

# ------------------------------------------------------------------
# 2. DECLARATIONS
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("2. DECLARATIONS")
print("=" * 100)

declarations = [
    "Funding",
    "Competing Interests",
    "Conflict of Interest",
    "Data and Code Availability",
    "Author Contributions",
    "CRediT",
    "Ethics Approval",
    "AI"
]

all_text = "\n".join(x.text for x in doc.paragraphs).lower()

for d in declarations:
    found = d.lower() in all_text
    print(f"{d}: {'PASS' if found else 'REVIEW'}")

# ------------------------------------------------------------------
# 3. STRUCTURE
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("3. MAIN STRUCTURE")
print("=" * 100)

structure_keywords = [
    "ABSTRACT",
    "1. INTRODUCTION",
    "2.",
    "3.",
    "4.",
    "5.",
    "6.",
    "7. RESEARCH LIMITATIONS AND FUTURE RESEARCH AGENDA",
    "CONCLUSION",
    "REFERENCES"
]

for s in structure_keywords:
    found = any(s.lower() in x.text.strip().lower() for x in doc.paragraphs)
    print(f"{s}: {'PASS' if found else 'REVIEW'}")

# ------------------------------------------------------------------
# 4. MALFORMED LATEX
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("4. LATEX / FORMULA CHECK")
print("=" * 100)

bad_latex = []

for i, para in enumerate(doc.paragraphs, 1):
    t = para.text

    # Valid forms are \text{in}, \text{out}, etc.
    # Detect malformed forms where the backslash is missing.
    if re.search(r'C_\{\s*ext\{(in|out)\}', t):
        bad_latex.append((i, t))

    if "C_{ ext" in t or "C_{      ext" in t:
        bad_latex.append((i, t))

if bad_latex:
    for i, t in bad_latex:
        print(f"BAD P{i}: {t}")
else:
    print("PASS — no malformed C_in/C_out LaTeX detected.")

# ------------------------------------------------------------------
# 5. SPACING / CONCATENATED WORDS
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("5. TARGETED SPACING CHECK")
print("=" * 100)

known_bad = [
    "thispattern",
    "byrecorded",
    "account,his",
    "thecommunicative",
    "theobserved",
    "crisiscommunication",
    "verifiedinvestigative",
    "network,however",
    "involvinginstitutional",
    "policiesare",
    "thisdynamic",
    "primaryarena",
    "darkirony",
    "positive'or",
    "andconsumer",
    "conductedindependently",
    "haveinappropriately",
    "facts.The",
    "capability,while",
    "separateTitle",
    "providedin",
    "forDouble",
    "CorporateReputation",
    "ofdeep",
    "inquiryinto",
    "ofSciences",
    "In2018",
]

spacing_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    t = para.text

    for bad in known_bad:
        if bad in t:
            spacing_hits.append((i, bad, t))

if spacing_hits:
    for i, bad, t in spacing_hits:
        print(f"P{i}: {bad}")
        print(t)
else:
    print("PASS — no known targeted spacing errors.")

# ------------------------------------------------------------------
# 6. STRONG / OVERINTERPRETIVE CLAIMS
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("6. STRONG CLAIM REVIEW")
print("=" * 100)

strong_terms = [
    "Algorithmic Oracle",
    "passive receptacles of grievance",
    "role abandoned by state public relations",
    "most transformative empirical finding",
    "dominating the graph's broadcasting capability",
    "remained silent",
    "absolute zero",
    "severe trauma",
    "rancid physical meals",
    "dismantle the conditions",
    "paralyze institutional legitimacy"
]

strong_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    t = para.text

    for term in strong_terms:
        if term.lower() in t.lower():
            strong_hits.append((i, term, t))

if strong_hits:
    for i, term, t in strong_hits:
        print(f"REVIEW P{i}: {term}")
        print(t)
else:
    print("PASS — no targeted high-risk wording detected.")

# ------------------------------------------------------------------
# 7. KEY EMPIRICAL NUMBERS
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("7. KEY EMPIRICAL VALUES")
print("=" * 100)

numbers = [
    "5,263",
    "3,395",
    "9.28%",
    "0.0011",
    "1.21%",
    "0.9837",
    "9.17%",
    "42",
    "15",
    "57.45%",
    "0.1444",
    "0.4563",
    "78.91%",
    "77.01%",
    "71.13%"
]

for n in numbers:
    print(f"{n}: {'FOUND' if n in all_text else 'REVIEW'}")

# ------------------------------------------------------------------
# 8. FORMATTING
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("8. FORMATTING")
print("=" * 100)

if doc.sections:
    s = doc.sections[0]

    print("Margins:")
    print("Top   :", round(s.top_margin.inches, 2))
    print("Bottom:", round(s.bottom_margin.inches, 2))
    print("Left  :", round(s.left_margin.inches, 2))
    print("Right :", round(s.right_margin.inches, 2))

for i in [1, 2, 3, 6, 75, 126, 158, 182, 184]:
    if i <= len(doc.paragraphs):
        p0 = doc.paragraphs[i-1]
        print(
            f"P{i}: style={p0.style.name}, "
            f"font={p0.runs[0].font.name if p0.runs else None}, "
            f"size={p0.runs[0].font.size.pt if p0.runs and p0.runs[0].font.size else None}, "
            f"spacing={p0.paragraph_format.line_spacing}"
        )

# ------------------------------------------------------------------
# 9. CRITICAL PARAGRAPHS
# ------------------------------------------------------------------
print("\n" + "=" * 100)
print("9. CRITICAL PARAGRAPHS")
print("=" * 100)

for i in [75, 76, 119, 126, 127, 158, 182, 184]:
    print(f"\nP{i}:")
    print(doc.paragraphs[i-1].text)

print("\n" + "=" * 100)
print("END OF FINAL HMC AUDIT V20")
print("=" * 100)
