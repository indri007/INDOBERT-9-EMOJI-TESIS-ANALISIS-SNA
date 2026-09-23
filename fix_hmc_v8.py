from docx import Document

src = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V7.docx"
out = "/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V8.docx"

doc = Document(src)

replacements = {
    "crisiscommunication": "crisis communication",
    "capability,while": "capability, while",
    "facts.The": "facts. The",
}

changed = 0

for i, para in enumerate(doc.paragraphs, 1):
    original = para.text
    updated = original

    for old, new in replacements.items():
        updated = updated.replace(old, new)

    if updated != original:
        if para.runs:
            para.runs[0].text = updated
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.add_run(updated)

        changed += 1
        print(f"CHANGED P{i}")

doc.save(out)

print()
print("CHANGED PARAGRAPHS:", changed)
print("OUTPUT:", out)
