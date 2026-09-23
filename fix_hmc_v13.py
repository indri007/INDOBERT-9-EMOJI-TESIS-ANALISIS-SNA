from docx import Document
from pathlib import Path

src = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V12.docx")
out = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V13.docx")

doc = Document(str(src))

replacements = {
    # ---------------------------------------------------------
    # Confirmed spacing corrections
    # ---------------------------------------------------------
    "policiesare": "policies are",
    "thisdynamic": "this dynamic",
    "Meal(Makan": "Meal (Makan",
    "In Indonesia,Platform X": "In Indonesia, Platform X",
    "stigmaby": "stigma by",
    "emojis(": "emojis (",
    "andportion": "and portion",
    "withpublic": "with public",
    "ofnetworked": "of networked",
    "networkedgatekeeping": "networked gatekeeping",
    "launchingpromotional": "launching promotional",
    "promotionaldigital": "promotional digital",
    "practicesin": "practices in",
    "fromcomplementing": "from complementing",
    "finalcorpus": "final corpus",
    "MBG)program": "MBG) program",
    "theanalyzed": "the analyzed",
    "haveinappropriately": "have inappropriately",
    "FAIRdata": "FAIR data",
    "separateTitle": "separate Title",
    "providedin": "provided in",
    "forDouble": "for Double",
    "CorporateReputation": "Corporate Reputation",
    "ofdeep": "of deep",
    "inquiryinto": "inquiry into",
    "ofSciences": "of Sciences",
    "103(23),8577": "103(23), 8577",
    "In2018": "In 2018",

    # ---------------------------------------------------------
    # LaTeX table correction
    # ---------------------------------------------------------
    r"C_{        ext{out}}": r"C_{\text{out}}",

    # ---------------------------------------------------------
    # Scholarly wording — P6
    # ---------------------------------------------------------
    "has fundamentally altered how public policies":
        "has substantially altered how public policies",

    # ---------------------------------------------------------
    # Scholarly wording — P8
    # ---------------------------------------------------------
    "inevitably misinterpret such sarcastic praise as genuine government support, severely misleading public policy assessments":
        "may misinterpret such sarcastic praise as genuine government support, potentially affecting the accuracy of public-policy assessments",
}

changed = []

for i, para in enumerate(doc.paragraphs, 1):
    old = para.text
    new = old

    for a, b in replacements.items():
        new = new.replace(a, b)

    if new != old:
        if para.runs:
            para.runs[0].text = new
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.add_run(new)

        changed.append((i, old, new))

# -------------------------------------------------------------
# P18: evidence-bounded rewrite
# -------------------------------------------------------------
for i, para in enumerate(doc.paragraphs, 1):
    if "compels networked citizens" in para.text:
        old = para.text

        new = old.replace(
            "Literature has yet to document or formalize how total institutional communicative paralysis "
            "($C_{\\text{out}} = 0$) compels networked citizens to abandon human authorities and summon "
            "AI conversational agents (@grok, $C_{\\text{out}} = 42$) as decentralized, third-party epistemic "
            "arbiters of state truth.",
            "The observed network provides an empirical basis for examining whether limited institutional "
            "outbound interaction ($C_{\\text{out}} = 0$) is associated with users seeking information from "
            "AI conversational agents such as @grok ($C_{\\text{out}} = 42$) as decentralized, third-party "
            "sources of policy-related information."
        )

        if new != old:
            if para.runs:
                para.runs[0].text = new
                for run in para.runs[1:]:
                    run.text = ""
            else:
                para.add_run(new)

            changed.append((i, old, new))

doc.save(str(out))

print("=" * 80)
print("HMC V13 CREATED")
print("=" * 80)
print("INPUT :", src)
print("OUTPUT:", out)
print("CHANGED PARAGRAPHS:", len(changed))
print()

for i, old, new in changed:
    print(f"P{i} CHANGED")
    print("OLD:", old[:250])
    print("NEW:", new[:250])
    print("-" * 80)

print("=" * 80)
