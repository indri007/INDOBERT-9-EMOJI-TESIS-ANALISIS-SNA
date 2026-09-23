from docx import Document
from pathlib import Path
import re

p = Path("/Users/jevin/Documents/tesis_mbg/hmc_submission/HMC_REVIEW_MANUSCRIPT_V12.docx")
doc = Document(str(p))

print("=" * 80)
print("FINAL HMC AUDIT — V12")
print("=" * 80)

# ------------------------------------------------------------
# 1. BASIC DOCUMENT
# ------------------------------------------------------------
print("\n[1] DOCUMENT")
print("Paragraphs :", len(doc.paragraphs))
print("Tables     :", len(doc.tables))
print("Sections   :", len(doc.sections))

# ------------------------------------------------------------
# 2. FIRST STRUCTURE
# ------------------------------------------------------------
print("\n[2] FIRST 12 PARAGRAPHS")

for i, para in enumerate(doc.paragraphs[:12], 1):
    text = para.text.strip()
    print(f"P{i}: {text[:200]}")

# ------------------------------------------------------------
# 3. DOUBLE-BLIND IDENTITY CHECK
# ------------------------------------------------------------
print("\n[3] DOUBLE-BLIND IDENTITY CHECK")

identity_terms = [
    "Indri Anjar Kartika Sari",
    "UPN Veteran Jawa Timur",
    "Universitas Pembangunan Nasional",
    "25067020011",
    "@student",
    "ORCID",
    "Corresponding author",
    "Email:",
    "Authors:",
    "Affiliation:",
]

identity_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    text = para.text
    for term in identity_terms:
        if term.lower() in text.lower():
            identity_hits.append((i, term, text[:180]))

if identity_hits:
    print("REVIEW — possible identity exposure:")
    for x in identity_hits:
        print(x)
else:
    print("PASS — no obvious author identity terms found.")

# ------------------------------------------------------------
# 4. METADATA / INTERNAL MARKERS
# ------------------------------------------------------------
print("\n[4] INTERNAL METADATA CHECK")

metadata_terms = [
    "Target Publication:",
    "Scopus ASJC Classification:",
    "Scopus Subject Area:",
    "SJR 0.76",
    "Names and Affiliations Anonymized",
]

metadata_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    text = para.text
    for term in metadata_terms:
        if term.lower() in text.lower():
            metadata_hits.append((i, term))

if metadata_hits:
    print("REVIEW:")
    for x in metadata_hits:
        print(x)
else:
    print("PASS")

# ------------------------------------------------------------
# 5. DECLARATIONS
# ------------------------------------------------------------
print("\n[5] DECLARATIONS")

declarations = [
    "Funding",
    "Competing Interests",
    "Conflict of Interest",
    "Data and Code Availability",
    "Author Contributions",
    "CRediT",
    "Ethics Approval",
    "AI",
]

full_text = "\n".join(p.text for p in doc.paragraphs).lower()

for d in declarations:
    print(f"{d:35}:", "PASS" if d.lower() in full_text else "REVIEW")

# ------------------------------------------------------------
# 6. MAIN STRUCTURE
# ------------------------------------------------------------
print("\n[6] MAIN STRUCTURE")

required_headings = [
    "ABSTRACT",
    "1. INTRODUCTION",
    "2. THEORETICAL FRAMEWORK AND LITERATURE REVIEW",
    "3. METHODOLOGY AND RESEARCH DESIGN",
    "4. EMPIRICAL RESULTS",
    "5. DISCUSSION",
    "6. THEORETICAL AND PRACTICAL CONTRIBUTIONS",
    "7. RESEARCH LIMITATIONS AND FUTURE RESEARCH AGENDA",
    "8. CONCLUSION",
    "STATEMENTS AND DECLARATIONS",
    "REFERENCES",
]

for h in required_headings:
    found = any(p.text.strip() == h for p in doc.paragraphs)
    print(f"{h:65}:", "PASS" if found else "REVIEW")

# ------------------------------------------------------------
# 7. KNOWN MALFORMED PATTERNS
# ------------------------------------------------------------
print("\n[7] MALFORMED TEXT CHECK")

bad_patterns = [
    r"aspect-basedsentiment",
    r"adedicated",
    r"incongruenceindicators",
    r"amodularity",
    r"categories\.The",
    r"networkstructure",
    r"policiesare",
    r"instead,policies",
    r"millionsof",
    r"thisdynamic",
    r"Meal\(Makan",
    r"civicdissent",
    r"publicsentiment",
    r"linguisticchallenges",
    r"withpublic",
    r"ofnetworked",
    r"networkedgatekeeping",
    r"launchingpromotional",
    r"promotionaldigital",
    r"practicesin",
    r"fromcomplementing",
    r"finalcorpus",
    r"MBG\)program",
    r"theanalyzed",
    r"FAIRdata",
    r"separateTitle",
    r"providedin",
    r"forDouble",
    r"CorporateReputation",
    r"ofdeep",
    r"inquiryinto",
    r"ofSciences",
    r"103\(23\),8577",
    r"In2018",
]

bad_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    text = para.text
    for pattern in bad_patterns:
        if re.search(pattern, text, re.IGNORECASE):
            bad_hits.append((i, pattern, text[:200]))

if bad_hits:
    print("REVIEW:")
    for x in bad_hits:
        print(x)
else:
    print("PASS — no known malformed patterns.")

# ------------------------------------------------------------
# 8. STRONG CLAIM FLAGS — REVIEW ONLY
# ------------------------------------------------------------
print("\n[8] STRONG CLAIM REVIEW")

strong_terms = [
    "compels",
    "abandon human authorities",
    "state truth",
    "severely misleading",
    "fundamentally altered",
    "inevitably misinterpret",
]

strong_hits = []

for i, para in enumerate(doc.paragraphs, 1):
    text = para.text
    for term in strong_terms:
        if term.lower() in text.lower():
            strong_hits.append((i, term, text[:240]))

if strong_hits:
    for x in strong_hits:
        print("REVIEW:", x)
else:
    print("PASS")

# ------------------------------------------------------------
# 9. TABLE / FIGURE MARKERS
# ------------------------------------------------------------
print("\n[9] TABLE / FIGURE CHECK")

for i, para in enumerate(doc.paragraphs, 1):
    text = para.text.strip()

    if (
        text.lower().startswith("table ")
        or text.lower().startswith("figure ")
        or text.startswith("|")
    ):
        print(f"P{i}: {text[:180]}")

# ------------------------------------------------------------
# 10. FORMATTING SAMPLE
# ------------------------------------------------------------
print("\n[10] FORMATTING SAMPLE")

for i in [1, 2, 3, 4, 5, 6, 7, 20, 50, 100, 150, 200]:
    if i <= len(doc.paragraphs):
        para = doc.paragraphs[i-1]

        print(
            f"P{i}: "
            f"style={para.style.name if para.style else None}; "
            f"line_spacing={para.paragraph_format.line_spacing}; "
            f"before={para.paragraph_format.space_before}; "
            f"after={para.paragraph_format.space_after}"
        )

# ------------------------------------------------------------
# 11. MARGINS
# ------------------------------------------------------------
print("\n[11] MARGINS")

for i, sec in enumerate(doc.sections, 1):
    print(
        f"Section {i}: "
        f"top={sec.top_margin.inches:.2f}, "
        f"bottom={sec.bottom_margin.inches:.2f}, "
        f"left={sec.left_margin.inches:.2f}, "
        f"right={sec.right_margin.inches:.2f}"
    )

print("\n" + "=" * 80)
print("END OF V12 AUDIT")
print("=" * 80)
