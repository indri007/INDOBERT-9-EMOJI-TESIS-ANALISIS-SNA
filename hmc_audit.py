from pathlib import Path
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
import re

# ============================================================
# HMC JOURNAL - PYTHON AUDIT
# ============================================================

SOURCE = Path(
    "/Users/jevin/Documents/tesis_mbg/hmc_submission/"
    "HMC_REVIEW_MANUSCRIPT.docx"
)

OUTPUT = Path(
    "/Users/jevin/Documents/tesis_mbg/hmc_submission/"
    "HMC_MANUSCRIPT_AUDIT_PY.txt"
)

if not SOURCE.exists():
    raise FileNotFoundError(f"MANUSCRIPT TIDAK DITEMUKAN:\n{SOURCE}")

doc = Document(SOURCE)

paragraphs = doc.paragraphs
all_text = "\n".join(p.text for p in paragraphs)

results = []

def status(label, value, detail=""):
    results.append({
        "label": label,
        "status": value,
        "detail": detail
    })

# ============================================================
# 1. FILE
# ============================================================

status(
    "DOCX FILE",
    "PASS",
    str(SOURCE)
)

# ============================================================
# 2. DOCUMENT METADATA / ANONYMIZATION
# ============================================================

core = doc.core_properties

metadata_fields = {
    "author": core.author,
    "last_modified_by": core.last_modified_by,
    "title": core.title,
    "subject": core.subject,
    "keywords": core.keywords,
    "comments": core.comments,
}

metadata_hits = {
    k: v for k, v in metadata_fields.items()
    if v
}

status(
    "DOCUMENT METADATA",
    "REVIEW" if metadata_hits else "PASS",
    str(metadata_hits) if metadata_hits else "No populated metadata fields."
)

# Author-identifying terms
identifying_terms = [
    "Indri Anjar Kartika Sari",
    "UPN VJ",
    "UPN Veteran Jawa Timur",
    "Universitas Pembangunan Nasional",
]

found_identifying = []

for term in identifying_terms:
    if term.lower() in all_text.lower():
        found_identifying.append(term)

status(
    "ANONYMIZATION",
    "FAIL" if found_identifying else "PASS",
    (
        "Found: " + ", ".join(found_identifying)
        if found_identifying
        else "Configured author-identifying terms not found."
    )
)

# ============================================================
# 3. ABSTRACT
# ============================================================

abstract_index = None

for i, p in enumerate(paragraphs):
    if p.text.strip().upper() == "ABSTRACT":
        abstract_index = i
        break

if abstract_index is None:
    status(
        "ABSTRACT",
        "FAIL",
        "ABSTRACT heading not found."
    )
else:
    abstract_parts = []

    for p in paragraphs[abstract_index + 1:]:
        txt = p.text.strip()

        if re.match(
            r"^1\.\s+INTRODUCTION$",
            txt,
            re.I
        ):
            break

        if txt:
            abstract_parts.append(txt)

    abstract_text = " ".join(abstract_parts)
    abstract_words = abstract_text.split()
    abstract_count = len(abstract_words)

    if 150 <= abstract_count <= 250:
        abstract_status = "PASS"
    else:
        abstract_status = "FAIL"

    status(
        "ABSTRACT 150-250 WORDS",
        abstract_status,
        f"{abstract_count} words"
    )

# ============================================================
# 4. KEYWORDS
# ============================================================

keyword_paragraph = None

for p in paragraphs:
    if re.match(
        r"^\s*Keywords?\s*:",
        p.text,
        re.I
    ):
        keyword_paragraph = p.text.strip()
        break

if not keyword_paragraph:
    status(
        "KEYWORDS 4-6",
        "FAIL",
        "Keywords line not found."
    )
else:
    keyword_text = re.sub(
        r"^\s*Keywords?\s*:\s*",
        "",
        keyword_paragraph,
        flags=re.I
    )

    # Accept semicolon, comma, or pipe separated keywords
    keywords = [
        x.strip()
        for x in re.split(
            r";|\||,",
            keyword_text
        )
        if x.strip()
    ]

    keyword_count = len(keywords)

    status(
        "KEYWORDS 4-6",
        "PASS" if 4 <= keyword_count <= 6 else "FAIL",
        f"{keyword_count}: {keywords}"
    )

# ============================================================
# 5. FONT
# ============================================================

font_errors = []

for i, p in enumerate(paragraphs, start=1):
    for r in p.runs:
        if not r.text.strip():
            continue

        size = r.font.size.pt if r.font.size else None
        name = r.font.name

        if size is not None and abs(size - 12) > 0.01:
            font_errors.append(
                f"Paragraph {i}: {size} pt"
            )

        if name and name.lower() not in [
            "times new roman",
            "timesnewroman",
        ]:
            font_errors.append(
                f"Paragraph {i}: font={name}"
            )

status(
    "12-POINT STANDARD FONT",
    "PASS" if not font_errors else "REVIEW",
    "No font-size/name exceptions detected."
    if not font_errors
    else "; ".join(font_errors[:20])
)

# ============================================================
# 6. DOUBLE SPACING
# ============================================================

spacing_errors = []

for i, p in enumerate(paragraphs, start=1):
    pf = p.paragraph_format

    if pf.line_spacing_rule is not None:
        if pf.line_spacing_rule != WD_LINE_SPACING.DOUBLE:
            spacing_errors.append(
                f"Paragraph {i}: {pf.line_spacing_rule}"
            )

status(
    "DOUBLE SPACING",
    "PASS" if not spacing_errors else "REVIEW",
    "No explicit non-double spacing detected."
    if not spacing_errors
    else "; ".join(spacing_errors[:20])
)

# ============================================================
# 7. MARGINS
# ============================================================

margin_errors = []

for i, section in enumerate(doc.sections, start=1):
    margins = {
        "top": section.top_margin,
        "bottom": section.bottom_margin,
        "left": section.left_margin,
        "right": section.right_margin,
    }

    for name, value in margins.items():
        if value is None:
            continue

        inches = value / 914400

        if abs(inches - 1.0) > 0.01:
            margin_errors.append(
                f"Section {i} {name}: {inches:.2f} inch"
            )

status(
    "1-INCH MARGINS",
    "PASS" if not margin_errors else "REVIEW",
    "All margins approximately 1 inch."
    if not margin_errors
    else "; ".join(margin_errors)
)

# ============================================================
# 8. PAGE NUMBERS
# ============================================================

page_field_found = False

for section in doc.sections:
    for p in section.footer.paragraphs:
        xml = p._p.xml
        if "PAGE" in xml:
            page_field_found = True

status(
    "PAGE NUMBERS",
    "PASS" if page_field_found else "REVIEW",
    "PAGE field detected in footer."
    if page_field_found
    else "No PAGE field detected."
)

# ============================================================
# 9. CONTINUOUS LINE NUMBERS
# ============================================================

settings_xml = doc.settings.element.xml

line_number_found = (
    "lnNumType" in settings_xml
)

status(
    "CONTINUOUS LINE NUMBERS",
    "PASS" if line_number_found else "REVIEW",
    "Line-numbering XML detected."
    if line_number_found
    else "Line-numbering XML not detected."
)

# ============================================================
# 10. UNDERLINE AUDIT
# ============================================================

underline_hits = []

for i, p in enumerate(paragraphs, start=1):
    for r in p.runs:
        if r.underline:
            underline_hits.append(
                f"Paragraph {i}: {r.text[:80]}"
            )

status(
    "UNDERLINE AUDIT",
    "REVIEW" if underline_hits else "PASS",
    "Underlined text found: " + str(underline_hits[:10])
    if underline_hits
    else "No underlined text detected."
)

# ============================================================
# 11. STRONG CLAIM AUDIT
# ============================================================

strong_claims = [
    "absolute breakdown",
    "empirical proof",
    "destroys democratic legitimacy",
    "textbook manifestation",
    "sociologically profound",
    "functionally dead",
    "historic paradigm shift",
    "cannot cure",
    "must abandon",
    "obsolete",
]

claim_hits = []

for phrase in strong_claims:
    matches = re.findall(
        re.escape(phrase),
        all_text,
        flags=re.I
    )

    if matches:
        claim_hits.append(phrase)

status(
    "STRONG CLAIM AUDIT",
    "REVIEW" if claim_hits else "PASS",
    (
        "Remaining phrases: " +
        ", ".join(claim_hits)
        if claim_hits
        else "No configured strong-claim phrases remain."
    )
)

# ============================================================
# 12. FIGURE CALLOUT AUDIT
# ============================================================

figure_numbers = sorted(set(
    re.findall(
        r"\bFigure\s+(\d+)",
        all_text,
        flags=re.I
    )
), key=int)

table_numbers = sorted(set(
    re.findall(
        r"\bTable\s+(\d+)",
        all_text,
        flags=re.I
    )
), key=int)

status(
    "FIGURE CALLOUTS",
    "INFO",
    f"Figures referenced: {figure_numbers}"
)

status(
    "TABLE CALLOUTS",
    "INFO",
    f"Tables referenced: {table_numbers}"
)

# ============================================================
# 13. EMBEDDED IMAGES
# ============================================================

inline_shapes = len(doc.inline_shapes)

status(
    "EMBEDDED FIGURES",
    "INFO",
    f"{inline_shapes} inline shape(s) detected."
)

# HMC review manuscripts should not depend on high-resolution
# production media. This script reports presence only.

# ============================================================
# 14. TABLE AUDIT
# ============================================================

table_count = len(doc.tables)

status(
    "TABLE OBJECTS",
    "INFO",
    f"{table_count} Word table object(s) detected."
)

# ============================================================
# 15. REFERENCES SECTION
# ============================================================

references_index = None

for i, p in enumerate(paragraphs):
    if p.text.strip().upper() == "REFERENCES":
        references_index = i
        break

if references_index is None:
    status(
        "REFERENCES SECTION",
        "FAIL",
        "REFERENCES heading not found."
    )
else:
    status(
        "REFERENCES SECTION",
        "PASS",
        f"Found at paragraph {references_index + 1}."
    )

# ============================================================
# 16. IN-TEXT CITATION AUDIT
# ============================================================

# Author-year citations such as:
# (Grice, 1975)
# (Kotler et al., 2023)
# (Newman, 2006; Papacharissi, 2015)

citation_pattern = re.compile(
    r"\(([A-Z][A-Za-zÀ-ÿ'’-]+)"
    r"(?:\s+et al\.)?"
    r"(?:,\s*|\s+)"
    r"\d{4}[a-z]?"
    r"(?:;\s*[^)]*)?\)"
)

citations = citation_pattern.findall(all_text)

status(
    "IN-TEXT CITATION DETECTION",
    "INFO",
    f"Approximate author-year citation groups detected: {len(citations)}"
)

# ============================================================
# 17. DOI AUDIT
# ============================================================

doi_pattern = re.compile(
    r"(?:https?://doi\.org/|doi:\s*)"
    r"10\.\d{4,9}/[-._;()/:A-Z0-9]+",
    re.I
)

doi_hits = doi_pattern.findall(all_text)

status(
    "DOI URL AUDIT",
    "INFO",
    f"DOI/DOI-URL occurrences detected: {len(doi_hits)}"
)

# ============================================================
# 18. REQUIRED HMC STATEMENTS
# ============================================================

required_statements = {
    "Funding": [
        r"\bFunding\b",
    ],
    "Conflict of Interest": [
        r"Conflict of Interest",
        r"Competing Interests",
    ],
    "Data Availability": [
        r"Data Availability",
        r"Data and Code Availability",
    ],
    "AI Disclosure": [
        r"AI Disclosure",
        r"AI / Tool Disclosure",
        r"Artificial Intelligence",
        r"generative AI",
        r"AI tools",
    ],
    "Ethics": [
        r"Ethics Approval",
        r"Ethical",
        r"Ethics Statement",
    ],
    "CRediT": [
        r"CRediT",
        r"Author Contributions",
    ],
}

for label, patterns in required_statements.items():
    found = any(
        re.search(pattern, all_text, re.I)
        for pattern in patterns
    )

    status(
        label,
        "PRESENT — AUTHOR VERIFICATION"
        if found
        else "MISSING / REVIEW",
        ""
    )

# ============================================================
# 19. DATA AVAILABILITY CONTENT CHECK
# ============================================================

data_terms = [
    "data availability",
    "dataset",
    "code availability",
    "repository",
    "github",
]

data_hits = [
    term for term in data_terms
    if term.lower() in all_text.lower()
]

status(
    "DATA AVAILABILITY CONTENT",
    "REVIEW",
    f"Detected terms: {data_hits}"
)

# ============================================================
# 20. AI DISCLOSURE CONTENT CHECK
# ============================================================

ai_terms = [
    "artificial intelligence",
    "generative ai",
    "chatgpt",
    "openai",
    "ai tool",
    "ai-assisted",
]

ai_hits = [
    term for term in ai_terms
    if term.lower() in all_text.lower()
]

status(
    "AI DISCLOSURE CONTENT",
    "REVIEW",
    f"Detected terms: {ai_hits}"
)

# ============================================================
# 21. ETHICS
# ============================================================

ethics_terms = [
    "ethics approval",
    "ethical approval",
    "ethics statement",
    "ethical considerations",
]

ethics_hits = [
    term for term in ethics_terms
    if term.lower() in all_text.lower()
]

status(
    "ETHICS CONTENT",
    "REVIEW",
    f"Detected terms: {ethics_hits}"
)

# ============================================================
# 22. CREDiT
# ============================================================

credit_hits = [
    term
    for term in [
        "CRediT",
        "Author Contributions",
        "Conceptualization",
        "Methodology",
        "Software",
        "Data Curation",
    ]
    if term.lower() in all_text.lower()
]

status(
    "CRediT CONTENT",
    "REVIEW",
    f"Detected terms: {credit_hits}"
)

# ============================================================
# 23. MARKDOWN RESIDUE
# ============================================================

markdown_patterns = [
    r"\*\*[^*]+\*\*",
    r"__[^_]+__",
    r"`[^`]+`",
]

markdown_hits = []

for pattern in markdown_patterns:
    if re.search(pattern, all_text):
        markdown_hits.append(pattern)

status(
    "MARKDOWN RESIDUE",
    "FAIL" if markdown_hits else "PASS",
    str(markdown_hits)
)

# ============================================================
# 24. REQUIRED MAIN SECTIONS
# ============================================================

main_sections = [
    "ABSTRACT",
    "1. INTRODUCTION",
    "2. THEORETICAL FRAMEWORK AND LITERATURE REVIEW",
    "3. METHODOLOGY AND RESEARCH DESIGN",
    "4. EMPIRICAL RESULTS",
    "5. DISCUSSION",
    "6. THEORETICAL AND PRACTICAL CONTRIBUTIONS",
    "7. RESEARCH LIMITATIONS AND FUTURE RESEARCH AGENDA",
    "8. CONCLUSION",
    "REFERENCES",
]

missing_sections = []

normalized = all_text.lower()

for section in main_sections:
    if section.lower() not in normalized:
        missing_sections.append(section)

status(
    "MAIN SECTION STRUCTURE",
    "PASS" if not missing_sections else "REVIEW",
    "Missing: " + ", ".join(missing_sections)
    if missing_sections
    else "Configured main sections detected."
)

# ============================================================
# 25. RESEARCH QUESTIONS
# ============================================================

rq_hits = []

for i in range(1, 7):
    if re.search(
        rf"\bRQ{i}\b",
        all_text,
        re.I
    ):
        rq_hits.append(f"RQ{i}")

status(
    "RESEARCH QUESTIONS",
    "PASS" if len(rq_hits) == 6 else "REVIEW",
    f"Detected: {rq_hits}"
)

# ============================================================
# 26. EMPIRICAL NUMBERS CONSISTENCY
# ============================================================

expected_values = [
    "5,310",
    "5,263",
    "3,395",
    "315",
    "9.28%",
    "971",
    "666",
    "692",
    "0.0011",
    "1.21%",
    "89",
    "9.17%",
    "332",
    "0.9837",
    "57.45%",
    "0.1444",
    "0.4563",
    "0.7178",
    "0.9692",
    "78.91%",
    "77.01%",
    "71.13%",
    "403",
    "535",
    "1,344",
    
]

missing_values = [
    value
    for value in expected_values
    if value not in all_text
]

status(
    "EMPIRICAL VALUE PRESENCE",
    "PASS" if not missing_values else "REVIEW",
    (
        "Missing configured values: " +
        ", ".join(missing_values)
        if missing_values
        else "Configured manuscript values detected."
    )
)

# ============================================================
# 27. FINAL SCORE
# ============================================================

hard_fail_labels = {
    "ABSTRACT 150-250 WORDS",
    "KEYWORDS 4-6",
    "ANONYMIZATION",
    "REFERENCES SECTION",
    "MARKDOWN RESIDUE",
}

fails = [
    x for x in results
    if x["label"] in hard_fail_labels
    and x["status"] == "FAIL"
]

passes = [
    x for x in results
    if x["status"] == "PASS"
]

reviews = [
    x for x in results
    if x["status"] in [
        "REVIEW",
        "PRESENT — AUTHOR VERIFICATION",
    ]
]

# This is NOT a journal acceptance probability/score.
# It is only a technical audit count.

# ============================================================
# 28. WRITE REPORT
# ============================================================

lines = []

lines.append("=" * 80)
lines.append("HMC JOURNAL — PYTHON MANUSCRIPT AUDIT")
lines.append("=" * 80)
lines.append("")
lines.append(f"Source : {SOURCE}")
lines.append(f"Output : {OUTPUT}")
lines.append("")

for item in results:
    lines.append(
        f"[{item['status']}] {item['label']}"
    )

    if item["detail"]:
        lines.append(
            f"    {item['detail']}"
        )

lines.append("")
lines.append("=" * 80)
lines.append("AUDIT SUMMARY")
lines.append("=" * 80)
lines.append(f"PASS   : {len(passes)}")
lines.append(f"REVIEW : {len(reviews)}")
lines.append(f"FAIL   : {len(fails)}")
lines.append("")

if fails:
    lines.append("CRITICAL FAILURES")
    lines.append("-" * 80)

    for item in fails:
        lines.append(
            f"- {item['label']}: {item['detail']}"
        )

else:
    lines.append("No configured hard-fail items detected.")

lines.append("")
lines.append("=" * 80)
lines.append("IMPORTANT")
lines.append("=" * 80)
lines.append(
    "This Python audit checks document structure and configured "
    "requirements. It does NOT certify journal acceptance."
)
lines.append(
    "APA 7 matching, figure placement, ethical compliance, "
    "AI disclosure accuracy, funding, conflict of interest, "
    "and data availability still require author verification."
)

OUTPUT.write_text(
    "\n".join(lines),
    encoding="utf-8"
)

print()
print("=" * 80)
print("HMC PYTHON AUDIT COMPLETE")
print("=" * 80)
print(f"SOURCE : {SOURCE}")
print(f"AUDIT  : {OUTPUT}")
print()
print(f"PASS   : {len(passes)}")
print(f"REVIEW : {len(reviews)}")
print(f"FAIL   : {len(fails)}")
print("=" * 80)
